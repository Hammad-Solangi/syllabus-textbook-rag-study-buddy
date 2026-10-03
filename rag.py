from __future__ import annotations

import io
import re

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np

from docx import Document as DocxDocument
from pypdf import PdfReader

from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
}


WORDS_PER_CHUNK = 220

OVERLAP_WORDS = 35


# =========================================================
# Data classes
# =========================================================

@dataclass(frozen=True)
class DocumentChunk:

    text: str

    source_name: str

    page: int | None

    chunk_number: int


    @property
    def source_label(self) -> str:

        if self.page is not None:

            return (
                f"{self.source_name} — "
                f"Page {self.page}"
            )

        return (
            f"{self.source_name} — "
            f"Document text"
        )


@dataclass(frozen=True)
class SearchHit:

    text: str

    source_name: str

    page: int | None

    chunk_number: int

    score: float


    @property
    def source_label(self) -> str:

        if self.page is not None:

            return (
                f"{self.source_name} — "
                f"Page {self.page}"
            )

        return (
            f"{self.source_name} — "
            f"Document text"
        )


# =========================================================
# Document index
# =========================================================

class DocumentIndex:

    def __init__(
        self,
        chunks: list[DocumentChunk],
    ):

        if not chunks:

            raise ValueError(
                "No readable text was found "
                "in the uploaded files."
            )


        self.chunks = chunks


        self.vectorizer = TfidfVectorizer(

            lowercase=True,

            strip_accents="unicode",

            ngram_range=(1, 2),

            sublinear_tf=True,

            min_df=1,

            max_features=50000,
        )


        self.matrix = (
            self.vectorizer.fit_transform(
                [chunk.text for chunk in chunks]
            )
        )


    # =====================================================
    # Create index from Streamlit uploads
    # =====================================================

    @classmethod
    def from_uploaded_files(
        cls,
        uploaded_files: Iterable,
    ) -> "DocumentIndex":

        chunks = []


        for uploaded in uploaded_files:

            name = uploaded.name

            suffix = Path(name).suffix.lower()

            raw = uploaded.getvalue()


            # ---------------------------------------------
            # PDF
            # ---------------------------------------------

            if suffix == ".pdf":

                pages = _extract_pdf_pages(raw)


                for page_number, text in pages:

                    new_chunks = _chunk_text(

                        text,

                        name,

                        page_number,

                        len(chunks) + 1,
                    )

                    chunks.extend(new_chunks)


            # ---------------------------------------------
            # DOCX
            # ---------------------------------------------

            elif suffix == ".docx":

                text = _extract_docx(raw)


                new_chunks = _chunk_text(

                    text,

                    name,

                    None,

                    len(chunks) + 1,
                )

                chunks.extend(new_chunks)


        return cls(chunks)


    # =====================================================
    # Document information
    # =====================================================

    @property
    def document_names(self) -> list[str]:

        return list(
            dict.fromkeys(
                chunk.source_name
                for chunk in self.chunks
            )
        )


    @property
    def document_count(self) -> int:

        return len(
            self.document_names
        )


    @property
    def chunk_count(self) -> int:

        return len(self.chunks)


    # =====================================================
    # Search
    # =====================================================

    def search(
        self,
        query: str,
        top_k: int = 5,
    ) -> list[SearchHit]:

        query = query.strip()


        if not query:

            return []


        query_vector = (
            self.vectorizer.transform(
                [query]
            )
        )


        scores = cosine_similarity(
            query_vector,
            self.matrix,
        ).ravel()


        ranked_indices = np.argsort(
            scores
        )[::-1]


        hits = []


        for index in ranked_indices:

            score = float(
                scores[index]
            )


            if score <= 0:

                continue


            chunk = self.chunks[
                int(index)
            ]


            hits.append(
                SearchHit(

                    text=chunk.text,

                    source_name=chunk.source_name,

                    page=chunk.page,

                    chunk_number=chunk.chunk_number,

                    score=score,
                )
            )


            if len(hits) >= top_k:

                break


        return hits


# =========================================================
# PDF extraction
# =========================================================

def _extract_pdf_pages(
    raw: bytes,
) -> list[tuple[int, str]]:

    reader = PdfReader(
        io.BytesIO(raw)
    )


    pages = []


    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):

        text = _clean_text(
            page.extract_text() or ""
        )


        if text:

            pages.append(
                (
                    page_number,
                    text,
                )
            )


    return pages


# =========================================================
# DOCX extraction
# =========================================================

def _extract_docx(
    raw: bytes,
) -> str:

    document = DocxDocument(
        io.BytesIO(raw)
    )


    parts = []


    # Normal paragraphs

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()


        if text:

            parts.append(text)


    # Tables

    for table in document.tables:

        for row in table.rows:

            cells = [
                cell.text.strip()
                for cell in row.cells
            ]


            row_text = " | ".join(
                cell
                for cell in cells
                if cell
            )


            if row_text:

                parts.append(
                    row_text
                )


    return _clean_text(
        "\n".join(parts)
    )


# =========================================================
# Chunking
# =========================================================

def _chunk_text(
    text: str,
    source_name: str,
    page: int | None,
    start_chunk: int,
) -> list[DocumentChunk]:

    text = _clean_text(text)


    words = text.split()


    if not words:

        return []


    chunks = []


    start = 0

    number = start_chunk


    while start < len(words):

        end = min(
            len(words),
            start + WORDS_PER_CHUNK,
        )


        chunk_text = " ".join(
            words[start:end]
        )


        chunks.append(
            DocumentChunk(

                text=chunk_text,

                source_name=source_name,

                page=page,

                chunk_number=number,
            )
        )


        number += 1


        if end >= len(words):

            break


        start = max(
            start + 1,
            end - OVERLAP_WORDS,
        )


    return chunks


# =========================================================
# Text cleaning
# =========================================================

def _clean_text(
    text: str,
) -> str:

    text = text.replace(
        "\x00",
        " ",
    )


    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )


    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )


    return text.strip()
